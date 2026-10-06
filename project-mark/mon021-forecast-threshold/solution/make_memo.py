import json
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm

f = json.load(open("facts.json")); S = {int(k): v for k, v in f["summary"].items()}
rep = pd.read_csv("replay_table.csv")
mon = lambda m: pd.Period(m).strftime("%B %Y")
short = lambda m: pd.Period(m).strftime("%b %Y")
flip = f["best_flip"]; bm = f["binding_month"]; b = S[flip]
r_bm = rep[(rep.threshold_pct == flip) & (rep.month == bm)].iloc[0]
miss100 = S[100]["missed_rows"][0]
m150 = pd.DataFrame(S[150]["missed_rows"], columns=["ba", "utc", "d", "dfc", "dimp", "dev"]).ba.value_counts()
v2bd = ", ".join(f"{k} {v}" for k, v in f["binding_breakdown_v2"].items())
m14bd = ", ".join(f"{k} {v}" for k, v in f["binding_breakdown_mon014"].items())
r50dec = rep[(rep.threshold_pct == 50) & (rep.month == "2025-12")].iloc[0]

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.7)


def para(text="", bold=False, size=None, space=4):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space); return p


def mixed(parts, space=4):
    p = doc.add_paragraph()
    for t, bb in parts:
        r = p.add_run(t); r.bold = bb
    p.paragraph_format.space_after = Pt(space); return p


def head(text):
    p = para(text, bold=True, size=11.5, space=2); p.paragraph_format.space_before = Pt(8)


def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet"); p.paragraph_format.space_after = Pt(2)


para("MON-021 v2 go-live decision (CHG-2291)", bold=True, size=15, space=2)
para("To: grid-data-ops on-call leads    From: grid-data-ops    Date: 6 October 2026", size=9.5, space=8)

head("Decision: hold. No candidate threshold can ship.")
mixed([("MON-021 v2 does not go live at any of the five candidate thresholds, so MON-021 stays off when v1 retires and "
        "CHG-2291 goes back to staffing review. ", True),
       (f"The binding month is {mon(bm)}. At {flip}%, the only candidate held back by nothing but capacity in a single "
        f"month, the primary on-call would have received {b['peak']} pages in {mon(bm)} against a triage capacity of "
        f"{b['peak_cap']}: {int(r_bm.mon021_v2_pages)} from MON-021 v2 and {int(r_bm.mon014_pages)} from MON-014, which "
        f"starts paging oncall-primary for the BAs outside v2's scope on the day v1 retires. ", False),
       (f"The smallest rota change that would let a candidate ship is {f['best_flip_total']} more triage pages in "
        f"{mon(bm)} (capacity {b['peak_cap']} to {b['peak']}); {flip}% would then pass both AC2 and AC3 and would be the "
        f"lowest passing candidate. No capacity change can rescue 100% or 150%, which fail coverage.", False)])

head("What blocks each candidate")
t = doc.add_table(rows=1, cols=5); t.style = "Light Grid Accent 1"
for i, h in enumerate(["Threshold", "Busiest month: on-call pages / capacity", "Months over capacity",
                       f"EIA-imputed hours missed (of {f['imputed_evaluated']})", "Blocked by"]):
    t.rows[0].cells[i].text = h
for th in [25, 50, 75, 100, 150]:
    s = S[th]
    rr = rep[(rep.threshold_pct == th) & (rep.month == s["peak_month"])].iloc[0]
    busiest = f"{short(s['peak_month'])}: {s['peak']} / {s['peak_cap']} (v2 {int(rr.mon021_v2_pages)} + MON-014 {int(rr.mon014_pages)})"
    blk = []
    if s["months_over"]:
        blk.append("capacity in every month" if s["months_over"] == 12 else
                   ("capacity in " + ", ".join(short(m) for m in s["over"])) if s["months_over"] <= 2 else
                   f"capacity in {s['months_over']} months, {short(s['over'][0])} to {short(s['over'][-1])}")
    if s["missed"]:
        blk.append(f"coverage ({s['missed']} missed)")
    row = t.add_row().cells
    for i, v in enumerate([f"{th}%", busiest, str(s["months_over"]), str(s["missed"]), "; ".join(blk)]):
        row[i].text = v
for row in t.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for r in p.runs: r.font.size = Pt(9)
para("", space=2)
bullet(f"50% first breaks in December 2025 ({int(r50dec.oncall_pages)} pages against a reduced capacity of "
       f"{int(r50dec.triage_capacity_pages)}) and stays over capacity in every month from then to September 2026, peaking at "
       f"{S[50]['peak']} in {mon(S[50]['peak_month'])}.")
bullet(f"100% misses one imputed hour: {miss100[0]}, hour ending {miss100[1][:16]} UTC, reported {float(miss100[2]):,.0f} MW "
       f"against a day-ahead forecast of {float(miss100[3]):,.0f} MW (EIA imputed {float(miss100[4]):,.0f} MW). The deviation "
       f"is {float(miss100[5]) * 100:.2f}%, just short of the 100% line. It is also one page over capacity in {mon(bm)} "
       f"({S[100]['peak']} vs {S[100]['peak_cap']}).")
bullet("150% fits the rota in every month but misses " + f"{S[150]['missed']} imputed hours (" +
       ", ".join(f"{k} {v}" for k, v in m150.items()) + "). 29 of the 34 are zero or negative demand values; "
       "every one deviates from forecast by between 99.9% and 132.9%, inside the 150% line.")

head(f"{mon(bm)} on-call pages at {flip}%, by BA")
mixed([(f"{b['peak']} pages. ", True),
       (f"MON-021 v2 ({int(r_bm.mon021_v2_pages)}): {v2bd}. MON-014 ({int(r_bm.mon014_pages)}): {m14bd}. ", False),
       ("May is MON-014's busiest month for the out-of-scope BAs, and that is what pushes it over: "
        "v2 alone stays inside capacity in every month at 75%.", False)])

head("Scope and the go-live paging set-up")
bullet(f"MON-021 v2 evaluates {f['evaluated_bas']} BAs. AVA, FPC, GVL, LGEE, PSEI, SEC, SPA and TEPC are outside its "
       "scope: EIA's EIA-930 documentation states that for these systems the day-ahead forecast includes commercial "
       "arrangements (pseudo-ties and dynamic scheduling) while reported demand is physical, so the comparison between "
       "actual and forecast demand is not very meaningful. Generation-only BAs report no demand or forecast and are never "
       "evaluated.")
bullet(f"Under the scheduled change in the Sentinel export, MON-014 pages oncall-primary for those eight BAs from v1 "
       f"retirement. Over the replay year it would have sent {f['mon014_total']} on-call pages ("
       + ", ".join(f"{k} {v}" for k, v in f["mon014_by_ba"].items()) +
       f"); {f['mon014_kinds'].get('floor', 0)} of those page-days are zero or negative demand only, "
       f"{f['mon014_kinds'].get('ceiling', 0)} are ceiling breaches only and {f['mon014_kinds'].get('ceiling,floor', 0)} have both.")
bullet(f"The back-of-envelope on the ticket is right about v2 on its own: at 75% it peaks at {S[75]['v2_alone_peak']} pages "
       f"in {mon(bm)} and never breaks capacity. It leaves out the MON-014 pages that join the same rota at go-live.")
bullet(f"For context, v1 sent {f['v1_total']:,} pages over the same twelve months, {f['v1_from_nc']:,} of them from the "
       f"eight out-of-scope BAs. The go-live set-up at 75% would have sent {S[75]['total']} "
       f"({S[75]['v2_total']} v2 + {f['mon014_total']} MON-014).")

head("How the replay was built")
bullet("Data: the five published EIA-930 BALANCE six-month files from July 2024 to October 2026, stacked on BA and "
       "'UTC Time at End of Hour' (no duplicate BA-hours). The two earliest files only feed MON-014's 365-day look-back. "
       "Pages are counted for UTC dates 1 October 2025 to 30 September 2026; the October 2026 rows in the live file are "
       "outside the window.")
bullet("MON-021 v2: reported 'Demand (MW)', not the adjusted series. An hour is evaluated when demand is reported and "
       "the forecast is above zero; it alerts when |demand - forecast| / forecast exceeds the threshold.")
bullet("MON-014: an hour is flagged when reported demand is zero or negative, or at least 1.5 times the highest "
       "unflagged demand in the previous 365 days. Flagged spikes such as SEC's 9,139,576 MW on 10 June 2026 are kept out "
       "of the ceiling history; letting them in would drop two pages (February and March 2026) without changing the decision.")
bullet("Paging and capacity: one page per BA per UTC date for each monitor, counted in that date's month. The rota's "
       "triage capacity covers every page routed to oncall-primary (24 in December 2025, 26 in July and August 2026, "
       "30 otherwise).")
bullet(f"Coverage: {f['imputed_evaluated']} hours that v2 evaluates have a populated 'Demand (MW) (Imputed)' value although "
       "the BA reported demand. Two more imputed FMPP hours had no forecast, so v2 cannot evaluate them and AC3 does not apply.")

head("Why no other conclusion holds")
para(f"Checking v2 against the rota on its own would ship 75%, but AC1 asks for the paging set-up that takes effect at "
     f"go-live, and from that day the eight out-of-scope BAs page the same analyst through MON-014. Running v2 across every "
     f"BA instead would also end in a hold, but for the wrong reason: forecasts that measure a different load would push "
     f"every candidate up to 100% over capacity in most months. With the scope and routing as configured, below 75% capacity "
     f"binds in most months, above 75% coverage fails, and 75% itself is {b['peak'] - b['peak_cap']} pages over in {mon(bm)} "
     f"alone. The call changes only if {mon(bm)} capacity rises to {b['peak']} or MON-014's routing for the out-of-scope BAs "
     f"is changed; the closest imputed hour sits at a {f['min_imputed_dev'] * 100:.2f}% deviation, so no 75% coverage "
     f"failure is close.", space=4)
doc.save("../golden/MON-021_v2_go-live_decision_memo.docx")
print("saved")
