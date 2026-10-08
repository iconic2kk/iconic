"""Builds the analysts' workbook template (scenario artifact). Its formulas follow Census's published rate convention
(average of the opening and closing July populations), which is NOT the agreement's start-of-year definition."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "inputs", "rrg_workbook_template.xlsx")
wb = Workbook(); inp = wb.active; inp.title = "Inputs"
blue = PatternFill("solid", fgColor="DDEBF7"); b = Font(bold=True)
inp["A1"] = "RRG evaluation workbook - template (board analysts, March 2026). Scenario artifact authored for this exercise."
inp["A2"] = "Paste Vintage 2025 values from co-est2025-alldata.csv into the blue cells. Rates are taken from, and formed like, the Census Bureau's published rate columns so they can be checked against the file."
cols = ["Area", "POPESTIMATE2022", "POPESTIMATE2023", "POPESTIMATE2024", "POPESTIMATE2025",
        "NATURALCHG2023", "NATURALCHG2024", "NATURALCHG2025", "DOMESTICMIG2023", "DOMESTICMIG2024", "DOMESTICMIG2025",
        "RNATURALCHG2023", "RNATURALCHG2024", "RNATURALCHG2025", "RDOMESTICMIG2023", "RDOMESTICMIG2024", "RDOMESTICMIG2025"]
for j, c in enumerate(cols, 1):
    inp.cell(4, j, c).font = b
inp["A5"] = "Dallas County (48113)"; inp["A6"] = "Texas (state row, SUMLEV 040)"
for r in (5, 6):
    for j in range(2, len(cols) + 1):
        inp.cell(r, j).fill = blue
inp.column_dimensions["A"].width = 30
# column letters: B..E pop22..pop25, F..H nat23..25, I..K dom23..25, L..N RNAT23..25, O..Q RDOM23..25
ev = wb.create_sheet("Evaluation")
for j, c in enumerate(["Area", "Year", "Natural change per 1,000", "Domestic migration per 1,000", "Domestically sourced growth per 1,000"], 1):
    ev.cell(1, j, c).font = b
years = [("Check year 2022-23", 0), ("Base year 2023-24", 1), ("Program year 2024-25", 2)]
nat, dom, rnat, rdom, pop = "FGH", "IJK", "LMN", "OPQ", "BCDE"
r = 2
for lab, k in years:   # Dallas: published rate columns
    ev.cell(r, 1, "Dallas County"); ev.cell(r, 2, lab)
    ev.cell(r, 3, f"=Inputs!{rnat[k]}5"); ev.cell(r, 4, f"=Inputs!{rdom[k]}5"); ev.cell(r, 5, f"=C{r}+D{r}"); r += 1
for lab, k in years:   # rest of Texas: state less Dallas, on the average of opening and closing populations
    den = f"(((Inputs!{pop[k]}6-Inputs!{pop[k]}5)+(Inputs!{pop[k+1]}6-Inputs!{pop[k+1]}5))/2)"
    ev.cell(r, 1, "Rest of Texas"); ev.cell(r, 2, lab)
    ev.cell(r, 3, f"=(Inputs!{nat[k]}6-Inputs!{nat[k]}5)/{den}*1000")
    ev.cell(r, 4, f"=(Inputs!{dom[k]}6-Inputs!{dom[k]}5)/{den}*1000"); ev.cell(r, 5, f"=C{r}+D{r}"); r += 1
ev["A10"] = "Effect (program year vs base year)"; ev["B10"] = "=(E4-E3)-(E7-E6)"
ev["A11"] = "Parallel-movement check (base year vs check year)"; ev["B11"] = "=(E3-E2)-(E6-E5)"
ev["A12"] = "Bar"; ev["B12"] = 1.8; ev["A13"] = "Limit"; ev["B13"] = 1.5
ev["A14"] = "Decision"; ev["B14"] = '=IF(AND(ABS(B11)<=B13,B10>=B12),"Extend","Lapse")'
for c in ("A10", "A11", "A14"): ev[c].font = b
ev.column_dimensions["A"].width = 46; ev.column_dimensions["B"].width = 22
for c in "CDE": ev.column_dimensions[c].width = 24
wb.save(OUT); print("saved", OUT)
