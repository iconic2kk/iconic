"""Builds the golden deliverables from inputs.zip: workbook (live formulas), board note (docx) and chart (png)."""
import sys, io, zipfile, json, os, subprocess, math
from fractions import Fraction as F
import pandas as pd
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
plt.rcParams['axes.unicode_minus'] = False
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from docx import Document
from docx.shared import Pt, Inches
HERE = os.path.dirname(os.path.abspath(__file__)); ZIP = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..', 'inputs.zip')
OUT = os.path.join(HERE, '..', 'golden'); os.makedirs(OUT, exist_ok=True)
subprocess.run([sys.executable, os.path.join(HERE, 'golden.py'), ZIP], check=True, capture_output=True)
g = json.load(open(os.path.join(HERE, 'figures.json')))
z = zipfile.ZipFile(ZIP)
q = pd.read_csv(io.BytesIO(z.read('qcew_aacog_counties_annual_2018_2025.csv')), dtype={'industry_code': str, 'area_fips': str})
ro = pd.read_csv(io.BytesIO(z.read('member_counties.csv')), dtype=str); nm = dict(zip(ro.county_fips, ro.county)); C = sorted(nm.values())
ti = pd.read_csv(io.BytesIO(z.read('qcew_industry_titles.csv')), dtype=str); title = dict(zip(ti.industry_code, ti.industry_title))
m = q[(q.own_code == 5) & (q.agglvl_code == 76) & q.industry_code.str[:2].isin(['31', '32', '33']) & (q.annual_avg_estabs > 0)].copy(); m['county'] = m.area_fips.map(nm)
P = lambda y: m[m.year == y].pivot_table(index='industry_code', columns='county', values='annual_avg_estabs', aggfunc='sum').reindex(columns=C).fillna(0).astype(int)
p24, p25 = P(2024), P(2025)
inds = sorted(set(p24.index) | set(p25.index))
d25 = {c: int(v) for c, v in g['draws']['2025'].items()}; d26 = g['draws_2026']; v2 = g['model_v2_2026']
pct = lambda a, b: (b - a) / a * 100
# exact fractional parts of 2026 draws (rounding robustness)
fr = {}
qual = [i for i in sorted(p25.index) if i in p24.index]
for c in C:
    x = sum(F(5_000_000, len(qual)) * F(int(p25.loc[i, c]), int(p25.loc[i].sum())) for i in qual)
    fr[c] = float(x - math.floor(x))
g['fraction_parts_2026'] = fr

# ---------- chart ----------
INK, INK2, MUTED, GRID, BASE, SURF = '#0b0b0b', '#52514e', '#898781', '#e1e0d9', '#c3c2b7', '#fcfcfb'
BLUE, ORANGE = '#2a78d6', '#eb6834'
order = sorted(C, key=lambda c: pct(d25[c], d26[c]))
fig, ax = plt.subplots(figsize=(7.6, 4.9), dpi=200); fig.patch.set_facecolor(SURF); ax.set_facecolor(SURF)
y = range(len(order))
for k, c in enumerate(order):
    a, b = pct(d25[c], d26[c]), pct(d25[c], v2[c])
    ax.plot([a, b], [k, k], color=GRID, lw=2, zorder=1)
ax.scatter([pct(d25[c], v2[c]) for c in order], y, s=46, color=ORANGE, edgecolor=SURF, linewidth=2, zorder=3, label="Finance office model v2 (projection deck)")
ax.scatter([pct(d25[c], d26[c]) for c in order], y, s=46, color=BLUE, edgecolor=SURF, linewidth=2, zorder=4, label='2026 round as the fund works it')
ax.axvline(-10, color=INK2, lw=1, ls=(0, (4, 3)), zorder=2)
ax.text(-9.6, len(order) - 0.6, 'Article 4 notice line (-10%)', ha='left', va='center', fontsize=7.5, color=INK2)
ax.axvline(0, color=BASE, lw=1, zorder=0)
ax.set_yticks(list(y)); ax.set_yticklabels(order, fontsize=8.5, color=INK2)
ax.set_xlabel('Change in draw, 2026 round against 2025 round (%)', fontsize=8.5, color=INK2)
ax.tick_params(axis='x', labelsize=8, colors=MUTED); ax.tick_params(axis='y', length=0)
ax.grid(axis='x', color=GRID, lw=0.6); ax.set_axisbelow(True)
for s in ax.spines.values(): s.set_visible(False)
w, gz = order.index('Wilson'), order.index('Gillespie')
ax.annotate(f"{pct(d25['Wilson'], d26['Wilson']):.1f}%", (pct(d25['Wilson'], d26['Wilson']), w), xytext=(0, 9), textcoords='offset points', fontsize=7.5, color=INK, ha='center')
ax.annotate(f"model v2 +{pct(d25['Wilson'], v2['Wilson']):.1f}%", (pct(d25['Wilson'], v2['Wilson']), w), xytext=(0, 9), textcoords='offset points', fontsize=7.5, color=INK2, ha='center')
ax.annotate(f"model v2 {pct(d25['Gillespie'], v2['Gillespie']):.1f}%", (pct(d25['Gillespie'], v2['Gillespie']), gz), xytext=(0, 9), textcoords='offset points', fontsize=7.5, color=INK2, ha='center')
ax.set_xlim(-16, 42); ax.set_ylim(-0.8, len(order) - 0.2)
ax.set_title('Only Wilson County falls past the notice line in the 2026 round', loc='left', fontsize=10.5, color=INK, pad=22)
ax.legend(loc='lower right', fontsize=7.5, frameon=False, bbox_to_anchor=(1.0, 1.0), ncol=2, handletextpad=0.3, columnspacing=1.2)
fig.tight_layout(); fig.savefig(os.path.join(OUT, 'draw_changes_2026.png'), facecolor=SURF); plt.close(fig)

# ---------- workbook ----------
wb = Workbook(); H = Font(bold=True); HF = PatternFill('solid', fgColor='DDE4EE')
s0 = wb.active; s0.title = 'Summary'
s1 = wb.create_sheet('Industries_2026'); s2 = wb.create_sheet('Draws_2026'); s3 = wb.create_sheet('Backtest_2020_2025')
s4 = wb.create_sheet('Model_v2_check'); s5 = wb.create_sheet('Medina_2025'); s6 = wb.create_sheet('Counts_2024_2025')
# Counts sheet: industry x county, 2024 and 2025
s6.append(['NAICS', 'Industry'] + [f'{c} 2024' for c in C] + ['Region 2024'] + [f'{c} 2025' for c in C] + ['Region 2025'])
for c in s6[1]: c.font = H; c.fill = HF
n = len(C)
for k, i in enumerate(inds, start=2):
    r24 = [int(p24.loc[i, c]) if i in p24.index else 0 for c in C]; r25 = [int(p25.loc[i, c]) if i in p25.index else 0 for c in C]
    s6.append([i, title.get(i, '')] + r24 + [f'=SUM(C{k}:{L(2 + n)}{k})'] + r25 + [f'=SUM({L(4 + n)}{k}:{L(3 + 2 * n)}{k})'])
NI = len(inds) + 1; reg24, reg25 = L(3 + n), L(4 + 2 * n)
# Industries sheet: qualification and part
s1.append(['NAICS', 'Industry', 'Region establishments 2024', 'Region establishments 2025', 'Carried in 2024 and 2025 (qualifies for 2026 round)', 'Part of the fund ($)'])
for c in s1[1]: c.font = H; c.fill = HF
for k, i in enumerate(inds, start=2):
    s1.append([i, title.get(i, ''), f'=Counts_2024_2025!{reg24}{k}', f'=Counts_2024_2025!{reg25}{k}', f'=IF(AND(C{k}>0,D{k}>0),1,0)', f'=IF(E{k}=1,5000000/SUM($E$2:$E${NI}),0)'])
s1.append([]); s1.append(['', 'Qualifying industries', '', '', f'=SUM(E2:E{NI})', f'=SUM(F2:F{NI})'])
# Draws sheet
s2.append(['County', 'Draw 2026 ($)', '2025 draw, ledger ($)', 'Change (%)', 'Article 4 notice', 'Transition payment ($)'])
for c in s2[1]: c.font = H; c.fill = HF
for k, c in enumerate(C, start=2):
    col = L(3 + n + C.index(c) + 1)  # county's 2025 column in Counts
    s2.append([c, f'=ROUND(SUMPRODUCT(Industries_2026!$F$2:$F${NI},Counts_2024_2025!{col}2:{col}{NI}/IF(Counts_2024_2025!${reg25}$2:${reg25}${NI}=0,1,Counts_2024_2025!${reg25}$2:${reg25}${NI})),0)',
               d25[c], f'=(B{k}-C{k})/C{k}*100', f'=IF(B{k}<0.9*C{k},"Yes","No")', f'=IF(E{k}="Yes",ROUND((C{k}-B{k})/2,0),0)'])
e = len(C) + 1
s2.append(['Total', f'=SUM(B2:B{e})', f'=SUM(C2:C{e})', '', f'=COUNTIF(E2:E{e},"Yes")', f'=SUM(F2:F{e})'])
for row in s2.iter_rows(min_row=2): row[3].number_format = '0.00'
# Backtest sheet (values from golden.py; ledger read from the pack)
led = pd.read_excel(io.BytesIO(z.read('breadth_fund_ledger_2020_2025.xlsx')), sheet_name='Draws', header=3)
s3.append(['Round', 'County', 'Ledger draw ($)', 'Rebuilt draw ($)', 'Difference', 'Qualifying industries', 'Part per industry ($)'])
for c in s3[1]: c.font = H; c.fill = HF
for k, r in enumerate(led.itertuples(index=False), start=2):
    R, cty, dl = int(r[0]), r[2], int(r[5])
    s3.append([R, cty, dl, g['draws'][str(R)][cty], f'=D{k}-C{k}', g['qualifying_industries'][str(R)], g['part_per_industry'][str(R)]])
s3.append([]); s3.append(['', 'Lines that differ', f'=COUNTIF(E2:E{len(led) + 1},"<>0")'])
# Model v2 sheet
s4.append(['County', 'Model v2 projection ($)', 'Draw 2026, fund method ($)', 'Overstatement ($)', 'Model v2 change (%)', 'Fund method change (%)'])
for c in s4[1]: c.font = H; c.fill = HF
for k, c in enumerate(C, start=2):
    s4.append([c, v2[c], f'=Draws_2026!B{k}', f'=B{k}-C{k}', round(pct(d25[c], v2[c]), 2), f'=Draws_2026!D{k}'])
s4.append([]); s4.append(['Industries model v2 counts that the region carried in 2025 but not in 2024 (first year, no part of the 2026 round):'])
for i in g['first_year_industries_2025']: s4.append([i, title.get(i, ''), ', '.join(f'{c}: {v} establishment(s)' for c, v in g['first_year_detail'][i].items())])
s4.append([]); s4.append(['Model v2 backtest, counties that differ from the ledger by round:'] + [f"{R}: {v}" for R, v in g['model_v2_backtest_mismatched_counties_by_round'].items()])
# Medina sheet
s5.append(['Industry', 'Title', 'Medina establishments (2023 counts)', 'Region (2023)', 'Medina part, 2024 round ($)', 'Medina establishments (2024 counts)', 'Region (2024)', 'Medina part, 2025 round ($)'])
for c in s5[1]: c.font = H; c.fill = HF
allm = sorted(set(g['medina_2024_parts']) | set(g['medina_2025_parts']))
for i in allm:
    a = g['medina_2024_parts'].get(i, [0, 0, 0]); b = g['medina_2025_parts'].get(i, [0, 0, 0])
    s5.append([i, title.get(i, ''), a[0], a[1], a[2], b[0], b[1], b[2]])
k = len(allm) + 1
s5.append(['Total', '', '', '', f'=SUM(E2:E{k})', '', '', f'=SUM(H2:H{k})'])
s5.append([]); s5.append(['June 2025 projection for Medina ($)', g['projection_2025']['Medina']]); s5.append(['Confirmed 2025 draw ($)', d25['Medina']]); s5.append(['Gap ($)', d25['Medina'] - g['projection_2025']['Medina']])
# Summary
for row in [['Breadth Fund 2026 round: golden summary'], [],
            ['Article 4 notices', ', '.join(g['notices_2026']) or 'none'], ['Transition reserve payout ($)', g['reserve_call_2026']],
            ['Wilson County 2026 draw ($)', d26['Wilson']], ['Wilson change on 2025 (%)', round(pct(d25['Wilson'], d26['Wilson']), 2)],
            ['Qualifying industries, 2026 round', g['qualifying_industries']['2026']], ['Part per industry ($)', g['part_per_industry']['2026']],
            ['Ledger lines rebuilt exactly (of 78)', 78 - g['backtest_mismatches']],
            ['Finance model v2: notices / reserve', f"{', '.join(g['model_v2_notices'])} / {g['model_v2_reserve_call']}"]]:
    s0.append(row)
s0['A1'].font = Font(bold=True, size=12)
for s in wb.worksheets:
    for col in range(1, min(s.max_column, 8) + 1): s.column_dimensions[L(col)].width = 22
x = os.path.join(OUT, 'breadth_fund_2026_round.xlsx'); wb.save(x)
subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', OUT + '/_recalc', x], capture_output=True)
os.replace(os.path.join(OUT, '_recalc', 'breadth_fund_2026_round.xlsx'), x); os.rmdir(os.path.join(OUT, '_recalc'))

# ---------- board note ----------
d = Document(); st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
d.add_heading('Breadth Fund 2026 round: notices, reserve and the Medina gap', level=1)
p = d.add_paragraph(); p.add_run('Notices and reserve. ').bold = True
p.add_run(f"One county gets an Article 4 notice: Wilson County, whose draw falls from ${d25['Wilson']:,} to ${d26['Wilson']:,} ({pct(d25['Wilson'], d26['Wilson']):.1f}%). "
          f"The transition reserve pays Wilson ${g['reserve_call_2026']:,}, half the fall. No other county falls by more than ten per cent; Gillespie is closest at {pct(d25['Gillespie'], d26['Gillespie']):.1f}%.")
p = d.add_paragraph(); p.add_run("Why the finance office's projection should not be issued. ").bold = True
p.add_run(f"Model v2 names Gillespie for the notice and a ${g['model_v2_reserve_call']:,} payment, and shows Wilson rising {pct(d25['Wilson'], v2['Wilson']):.1f}%. "
          "It reproduces the 2023 to 2025 rounds, but not 2020 to 2022: the fund backs an industry only once the region has carried it two years running, and model v2 counts every industry carried in the latest year. "
          "Two industries are in their first year in the 2025 counts, pulp and paper mills (one establishment in Bexar) and household appliances (two in Wilson), so they take no part in the 2026 round. Counting them hands Wilson a whole industry's part it is not yet due.")
p = d.add_paragraph(); p.add_run('What happened to Medina. ').bold = True
p.add_run(f"The fund splits each round equally across the industries the region has carried two years running (82 in the 2025 round, ${g['part_per_industry']['2025']:,.2f} each), and each industry's part follows the counties' shares of its establishments. "
          f"Medina had the region's only yarn and fiber mill (NAICS 3131), so it held that industry's whole part, ${g['medina_2025_parts'].get('3131', g['medina_2024_parts']['3131'])[2]:,.2f}, in the 2024 round. The mill is gone from the 2024 counts, the industry left the region, and Medina's draw fell to ${d25['Medina']:,}. "
          f"The June projection moved Medina's 2024 draw with its own establishment count (26 to 25) and so showed ${g['projection_2025']['Medina']:,}. The confirmed round was right; the projection method was the cause of the gap.")
d.add_picture(os.path.join(OUT, 'draw_changes_2026.png'), width=Inches(6.2))
t = d.add_table(rows=1, cols=4); t.style = 'Light Grid Accent 1'
for k, h in enumerate(['County', '2025 draw', '2026 draw', 'Change']): t.rows[0].cells[k].text = h
for c in C:
    r = t.add_row().cells; r[0].text = c; r[1].text = f'${d25[c]:,}'; r[2].text = f'${d26[c]:,}'; r[3].text = f'{pct(d25[c], d26[c]):+.1f}%'
d.add_paragraph('Every figure is rebuilt from the BLS annual average private establishment counts in the pack; the same method reproduces all 78 confirmed draws in the round ledger from 2020 to 2025 to the dollar.')
d.save(os.path.join(OUT, 'breadth_fund_2026_note.docx'))
json.dump(g, open(os.path.join(HERE, 'figures.json'), 'w'), indent=1)
print('fraction parts (2026):', {c: round(v, 3) for c, v in fr.items()})
