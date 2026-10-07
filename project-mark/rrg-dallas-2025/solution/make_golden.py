"""Builds the golden deliverables (workbook, note, chart) from the shipped inputs only."""
import os
import subprocess

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
IN = os.path.join(HERE, "..", "inputs")
OUT = os.path.join(HERE, "..", "golden")
os.makedirs(OUT, exist_ok=True)

d = pd.read_csv(os.path.join(IN, "co-est2025-alldata.csv"), encoding="latin-1")
tx = d[(d.STATE == 48) & (d.SUMLEV == 50)]
groups = {"Dallas County": tx[tx.CTYNAME == "Dallas County"].sum(numeric_only=True),
          "Rest of Texas": tx[tx.CTYNAME != "Dallas County"].sum(numeric_only=True)}
YEARS = [("Check year", 2023), ("Base year", 2024), ("Program year", 2025)]
BAR, BAND = 1.80, 1.50

rows = []
for area, s in groups.items():
    for lab, y in YEARS:
        base = s[f"POPESTIMATE{y-1}"]
        rows.append(dict(area=area, year=lab, period=f"1 Jul {y-1} to 1 Jul {y}", start_pop=int(base),
                         natural=int(s[f"NATURALCHG{y}"]), domestic=int(s[f"DOMESTICMIG{y}"]),
                         nat_rate=s[f"NATURALCHG{y}"] / base * 1000, dom_rate=s[f"DOMESTICMIG{y}"] / base * 1000))
t = pd.DataFrame(rows)
t["rate"] = t.nat_rate + t.dom_rate
R = t.set_index(["area", "year"])
def chg(area, a, b, col="rate"): return R.loc[(area, b), col] - R.loc[(area, a), col]
effect = chg("Dallas County", "Base year", "Program year") - chg("Rest of Texas", "Base year", "Program year")
check = chg("Dallas County", "Check year", "Base year") - chg("Rest of Texas", "Check year", "Base year")
extend = abs(check) <= BAND and effect >= BAR

# context: Census mid-year basis (what the published rate columns use) and total growth
def alt(area, y, kind):
    s = groups[area]; st = s[f"POPESTIMATE{y-1}"]; mid = (st + s[f"POPESTIMATE{y}"]) / 2
    nd = s[f"NATURALCHG{y}"] + s[f"DOMESTICMIG{y}"]
    return {"mid": nd / mid, "total": s[f"NPOPCHG{y}"] / st}[kind] * 1000
def alt_did(kind, y0, y1):
    return (alt("Dallas County", y1, kind) - alt("Dallas County", y0, kind)) - (alt("Rest of Texas", y1, kind) - alt("Rest of Texas", y0, kind))
mid_eff, tot_eff = alt_did("mid", 2024, 2025), alt_did("total", 2024, 2025)
dal = groups["Dallas County"]

# ---------- chart ----------
BLUE, ORANGE, INK, INK2, GRID, SURF = "#2a78d6", "#eb6834", "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"
fig, ax = plt.subplots(figsize=(6.6, 3.6), dpi=200)
fig.patch.set_facecolor(SURF); ax.set_facecolor(SURF)
x = [0, 1, 2]
for area, col, mk in [("Rest of Texas", ORANGE, "s"), ("Dallas County", BLUE, "o")]:
    v = [R.loc[(area, lab), "rate"] for lab, _ in YEARS]
    ax.plot(x, v, color=col, lw=2, marker=mk, ms=7, mec=SURF, mew=1.5, label=area, zorder=3)
    ax.annotate(f"{area}\n{v[-1]:+.2f}", (2, v[-1]), xytext=(10, 0), textcoords="offset points", va="center",
                fontsize=8, color=INK)
ax.axhline(0, color=INK2, lw=0.8, zorder=1)
ax.set_xticks(x, [f"{lab}\n{y-1}-{str(y)[2:]}" for lab, y in YEARS], fontsize=8, color=INK2)
ax.tick_params(axis="y", labelsize=8, colors=INK2, length=0); ax.tick_params(axis="x", length=0)
ax.set_ylabel("per 1,000 residents at start of year", fontsize=8, color=INK2)
ax.grid(axis="y", color=GRID, lw=0.6); ax.set_axisbelow(True)
for sp in ax.spines.values(): sp.set_visible(False)
ax.set_xlim(-0.25, 2.75); ax.set_ylim(-14, 16)
ax.set_title("Domestically sourced growth (natural change + net domestic migration)", fontsize=9.5, color=INK, loc="left")
ax.legend(loc="upper center", ncol=2, frameon=False, fontsize=8, bbox_to_anchor=(0.45, 1.0))
fig.tight_layout()
chart = os.path.join(OUT, "rrg_rates_chart.png"); fig.savefig(chart, facecolor=SURF); plt.close(fig)

# ---------- workbook ----------
wb = Workbook(); ws = wb.active; ws.title = "Evaluation"
hdr = Font(bold=True); fill = PatternFill("solid", fgColor="EEF3FB")
ws["A1"] = "RRG evaluation on Census Vintage 2025 (co-est2025-alldata): domestically sourced growth per 1,000 residents at the start of each estimate year"
ws["A1"].font = Font(bold=True, size=12)
cols = ["Area", "Year", "Estimate year", "Start-of-year population", "Natural change", "Net domestic migration",
        "Natural change per 1,000", "Domestic migration per 1,000", "Domestically sourced growth per 1,000"]
for j, c in enumerate(cols, 1):
    cell = ws.cell(row=3, column=j, value=c); cell.font = hdr; cell.fill = fill; cell.alignment = Alignment(wrap_text=True)
r0 = 4
for i, rw in t.iterrows():
    r = r0 + i
    ws.cell(r, 1, rw.area); ws.cell(r, 2, rw.year); ws.cell(r, 3, rw.period)
    ws.cell(r, 4, rw.start_pop); ws.cell(r, 5, rw.natural); ws.cell(r, 6, rw.domestic)
    ws.cell(r, 7, f"=E{r}/D{r}*1000"); ws.cell(r, 8, f"=F{r}/D{r}*1000"); ws.cell(r, 9, f"=G{r}+H{r}")
    for c in (7, 8, 9): ws.cell(r, c).number_format = "0.000"
    for c in (4, 5, 6): ws.cell(r, c).number_format = "#,##0"
# rows: Dallas 4,5,6 ; Rest 7,8,9
ws["A12"] = "Effect (program year vs base year)"; ws["A12"].font = hdr
ws["B12"] = "=(I6-I5)-(I9-I8)"; ws["C12"] = "points per 1,000"
ws["A13"] = "  of which natural change"; ws["B13"] = "=(G6-G5)-(G9-G8)"
ws["A14"] = "  of which domestic migration"; ws["B14"] = "=(H6-H5)-(H9-H8)"
ws["A15"] = "Bar (section 7)"; ws["B15"] = BAR
ws["A16"] = "Effect less bar"; ws["B16"] = "=B12-B15"
ws["A18"] = "Parallel-movement check (base year vs check year)"; ws["A18"].font = hdr
ws["B18"] = "=(I5-I4)-(I8-I7)"
ws["A19"] = "Limit (section 6)"; ws["B19"] = BAND
ws["A20"] = "Inside the limit by"; ws["B20"] = "=B19-ABS(B18)"
ws["A22"] = "Decision (section 7)"; ws["A22"].font = hdr
ws["B22"] = '=IF(AND(ABS(B18)<=B19,B12>=B15),"Extend for a second term (1 Jul 2027 to 30 Jun 2030)","Lapse at end of first term (30 Jun 2027)")'
ws["A24"] = "Annex A reproduction (check year, one decimal)"; ws["A24"].font = hdr
ws["A25"] = "Dallas County"; ws["B25"] = "=ROUND(I4,1)"; ws["C25"] = -6.2
ws["A26"] = "Rest of Texas"; ws["B26"] = "=ROUND(I7,1)"; ws["C26"] = 13.2
ws["A27"] = "Gap"; ws["B27"] = "=ROUND(I7-I4,1)"; ws["C27"] = 19.5
ws["D24"] = "Annex A"; ws["C24"] = "Annex A"; ws["B24"] = "Compiled"
for c in ("B12", "B13", "B14", "B16", "B18", "B20"): ws[c].number_format = "0.000"
ws.column_dimensions["A"].width = 44
for c in "BCDEFGHI": ws.column_dimensions[c].width = 17
ws2 = wb.create_sheet("Notes")
notes = [
    "Source: U.S. Census Bureau, Vintage 2025 county population estimates, co-est2025-alldata.csv, as published. All three years from this one release.",
    "Rest of Texas: the 253 Texas counties other than Dallas County (48113), counts summed before the rate is formed (equals the Texas state row less Dallas County).",
    "Numerator: NATURALCHG + DOMESTICMIG for the estimate year. International migration (INTERNATIONALMIG) and the Census residual (RESIDUAL) are not part of the board's measure.",
    "Denominator: the 1 July resident population that opens the estimate year (POPESTIMATE of the previous year). This is the only basis that reproduces Annex A (-6.2, +13.2, 19.5).",
    f"Census's published rate columns (RNATURALCHG, RDOMESTICMIG) divide by the average of the opening and closing July populations; on that basis the effect is {mid_eff:.3f}, below the bar. It does not reproduce Annex A and is not the board's calculation.",
    f"Total population change, which includes international migration, gives an effect of {tot_eff:.3f}; the agreement excludes international migration.",
]
for i, n in enumerate(notes, 1): ws2.cell(i, 1, n)
ws2.column_dimensions["A"].width = 160
xlsx = os.path.join(OUT, "rrg_evaluation_2025.xlsx"); wb.save(xlsx)
# recalculate so the formula results are cached in the file
tmp = os.path.join(OUT, "_recalc"); os.makedirs(tmp, exist_ok=True)
subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, xlsx], check=True, capture_output=True)
os.replace(os.path.join(tmp, "rrg_evaluation_2025.xlsx"), xlsx); os.rmdir(tmp)

# ---------- note ----------
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
doc.add_heading("Relocation and Retention Grant: extension decision", level=1)
doc.add_paragraph("To: County Workforce Board members    From: RRG evaluation    Meeting: 22 October 2026")
p = doc.add_paragraph()
p.add_run("Recommendation: extend the RRG for a second term, 1 July 2027 to 30 June 2030. ").bold = True
p.add_run(f"Both tests in the evaluation agreement are met, narrowly. The effect is +{effect:.3f} points per 1,000, "
          f"{effect-BAR:.3f} above the 1.80 bar. The parallel-movement check is {check:.3f}, inside the 1.50 limit by "
          f"{BAND-abs(check):.3f}. Both are compared before rounding, as section 7 requires.")
doc.add_heading("The board's measure, year by year", level=2)
tb = doc.add_table(rows=1, cols=4); tb.style = "Light Grid Accent 1"
for j, h in enumerate(["Per 1,000 residents at start of year", "Check year 2022-23", "Base year 2023-24", "Program year 2024-25"]):
    tb.rows[0].cells[j].text = h
for area in groups:
    c = tb.add_row().cells; c[0].text = area
    for j, (lab, _) in enumerate(YEARS, 1):
        c[j].text = f"{R.loc[(area, lab), 'rate']:+.2f}"
doc.add_paragraph(f"Dallas County moved by {chg('Dallas County','Base year','Program year'):+.3f} points from the base year to the "
                  f"program year; the rest of Texas moved by {chg('Rest of Texas','Base year','Program year'):+.3f}. The difference is the effect.")
doc.add_heading("Where the change came from", level=2)
for area in groups:
    dn = chg(area, "Base year", "Program year", "nat_rate"); dd = chg(area, "Base year", "Program year", "dom_rate")
    doc.add_paragraph(f"{area}: natural change {dn:+.3f}, domestic migration {dd:+.3f}, together {dn+dd:+.3f} points per 1,000.", style="List Bullet")
en = chg("Dallas County", "Base year", "Program year", "nat_rate") - chg("Rest of Texas", "Base year", "Program year", "nat_rate")
ed = chg("Dallas County", "Base year", "Program year", "dom_rate") - chg("Rest of Texas", "Base year", "Program year", "dom_rate")
doc.add_paragraph(f"Of the {effect:.3f} effect, domestic migration contributes {ed:+.3f} and natural change {en:+.3f}. Dallas County's net "
                  f"domestic out-migration eased from {abs(R.loc[('Dallas County','Base year'),'domestic']):,} to "
                  f"{abs(R.loc[('Dallas County','Program year'),'domestic']):,} people while the rest of Texas's domestic gain slowed.")
doc.add_picture(chart, width=Inches(6.2)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading("Points for members", level=2)
doc.add_paragraph(f"Did Dallas County's growth collapse? Total population change fell from {int(dal['NPOPCHG2024']):+,} to "
                  f"{int(dal['NPOPCHG2025']):+,}, almost entirely because net international migration dropped from "
                  f"{int(dal['INTERNATIONALMIG2024']):,} to {int(dal['INTERNATIONALMIG2025']):,}. The agreement leaves international "
                  f"migration out; on total growth the comparison would read {tot_eff:.2f}, which is not the board's question.", style="List Bullet")
doc.add_paragraph("Rates are formed on the resident population at the start of each estimate year, the basis of the board's own "
                  "calculation in Annex A: it returns Dallas -6.2, rest of Texas +13.2 and a gap of 19.5 for the check year. "
                  f"Census's published rate columns use the year's average population; on that basis the effect would be {mid_eff:.3f} "
                  "and the program would lapse, but that basis does not reproduce Annex A.", style="List Bullet")
doc.add_paragraph("The rest of Texas is the other 253 counties taken as one area (counts summed, then the rate formed). "
                  "The Census residual is not a component of change and is left out, as is international migration.", style="List Bullet")
doc.add_paragraph("Census's release notes say the 2024-25 domestic migration estimates used incomplete inputs and will be revised "
                  "in Vintage 2026. The agreement fixes Vintage 2025 as published, and with a 0.086 margin a later revision could "
                  "land either side of the bar.", style="List Bullet")
doc.save(os.path.join(OUT, "rrg_decision_note.docx"))
print(f"effect {effect:.4f} check {check:.4f} extend {extend} mid {mid_eff:.4f} total {tot_eff:.4f} nat {en:.4f} dom {ed:.4f}")
