"""Independent check in plain Python (csv module, integer arithmetic), sharing no code with golden.py."""
import csv, zipfile, io, sys, json
from openpyxl import load_workbook
zf = zipfile.ZipFile(sys.argv[1])
fips = {r['county_fips']: r['county'] for r in csv.DictReader(io.TextIOWrapper(zf.open('member_counties.csv')))}
cnt = {}   # (year, naics4) -> {county: estabs}
for r in csv.DictReader(io.TextIOWrapper(zf.open('qcew_aacog_counties_annual_2018_2025.csv'))):
    if r['own_code'] == '5' and r['agglvl_code'] == '76' and r['industry_code'][:2] in ('31', '32', '33') and int(r['annual_avg_estabs']) > 0:
        cnt.setdefault((int(r['year']), r['industry_code']), {})[fips[r['area_fips']]] = int(r['annual_avg_estabs'])
counties = sorted(fips.values())
def run(R):
    y = R - 1
    ok = sorted(i for (yy, i) in cnt if yy == y and (y - 1, i) in cnt)
    n = len(ok); num = {c: 0 for c in counties}; den = {}
    # draw_c = 5e6/n * sum_i e_ci/E_i ; exact with a common denominator per industry
    from math import lcm
    L = 1
    for i in ok: L = lcm(L, sum(cnt[(y, i)].values()))
    for i in ok:
        t = sum(cnt[(y, i)].values())
        for c, v in cnt[(y, i)].items(): num[c] += v * (L // t)
    # draw = 5_000_000 * num / (n * L), rounded half up
    return {c: (2 * 5_000_000 * num[c] + n * L) // (2 * n * L) for c in counties}
wb = load_workbook(io.BytesIO(zf.read('breadth_fund_ledger_2020_2025.xlsx')), data_only=True)
ws = wb['Draws']; bad = 0; lines = 0
for row in ws.iter_rows(min_row=5, values_only=True):
    if row[0] is None: continue
    lines += 1; bad += run(int(row[0]))[row[2]] != int(row[5])
d25, d26 = run(2025), run(2026)
notices = {c: (d25[c] - d26[c] + 1) // 2 for c in counties if 10 * d26[c] < 9 * d25[c]}
print(json.dumps({'ledger_lines': lines, 'mismatches': bad, 'draws_2026': d26, 'notices': notices, 'reserve': sum(notices.values())}))
