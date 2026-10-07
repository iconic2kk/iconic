import json, os
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, Inches
from scorecard import load, scored, error_pct, flags, INPUTS

b = flags(load()); s = scored(b)
m = error_pct(s, "month"); q = error_pct(s, "quarter")
g = pd.read_csv("golden_q3.csv", index_col=0)
out = pd.DataFrame({"ba": g.index, "q3_forecast_error_pct": g.q3_raw.round(1).values,
                    "jul_2026_pct": g["2026-07"].round(1).values, "aug_2026_pct": g["2026-08"].round(1).values,
                    "sep_2026_pct": g["2026-09"].round(1).values, "q3_rank": g["rank"].values,
                    "forecast_watch": g.watch.map({True: "Y", False: "N"}).values})
out["formal_letter"] = (out.q3_rank == 1).map({True: "Y", False: "N"})
out.to_csv("../golden/scorecard_q3_2026.csv", index=False)

# controls comparison
hist = pd.read_csv(os.path.join(INPUTS, "forecast_scorecard_published_2024Q3_2026Q2.csv"), skiprows=1)
hist["compiled"] = [round(float((m if r.period_type == "month" else q).loc[(r.ba, r.period)]), 1) for r in hist.itertuples()]
hist["difference"] = (hist.compiled - hist.forecast_error_pct).round(1)
assert (hist.difference == 0).all()
var = json.load(open("variants.json"))

# chart
fig, ax = plt.subplots(figsize=(10, 5.2))
cols = ["#c0392b" if w else "#8a9bb0" for w in g.watch]
ax.bar(range(len(g)), g.q3_raw, color=cols)
for i, v in enumerate(g.q3_raw):
    ax.text(i, v + 0.3, f"{v:.1f}", ha="center", fontsize=7.5)
ax.axvline(4.5, color="black", linestyle="--", linewidth=1.2)
ax.text(4.6, 24, "Forecast Watch cut\n(5th LDWP 8.65, 6th CISO 8.40)", fontsize=8.5, va="top")
ax.set_xticks(range(len(g))); ax.set_xticklabels(g.index, rotation=60, fontsize=8)
ax.set_ylabel("Q3 2026 forecast error, %")
ax.set_title("Q3 2026 forecast error by scorecard BA: PSCO, WALC, SWPW, SRP and LDWP go on Forecast Watch", fontsize=11, fontweight="bold")
plt.tight_layout(); plt.savefig("q3_chart.png", dpi=160); plt.close()

doc = Document(); st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = sec.bottom_margin = Cm(1.6)
def para(t, bold=False, size=None, space=4):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold
    if size: r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space); return p
def bullet(t):
    p = doc.add_paragraph(t, style="List Bullet"); p.paragraph_format.space_after = Pt(2)
top = g.head(6)
para("Q3 2026 forecast reliability scorecard and Q4 Forecast Watch", bold=True, size=14, space=2)
para("To: grid data desk leads    From: grid data desk    Date: 6 October 2026", size=9, space=8)
para("Decision", bold=True, size=11.5, space=2)
para(f"PSCO, WALC, SWPW, SRP and LDWP go on Forecast Watch for Q4 2026, and PSCO receives the formal letter. PSCO's Q3 forecast "
     f"error is {top.q3_raw.iloc[0]:.1f}%, {top.q3_raw.iloc[0] - top.q3_raw.iloc[1]:.2f} points above WALC at {top.q3_raw.iloc[1]:.1f}%. "
     f"The Watch cut falls between LDWP in fifth place at {top.q3_raw.iloc[4]:.2f}% and CISO in sixth at {top.q3_raw.iloc[5]:.2f}%, "
     f"a margin of {top.q3_raw.iloc[4] - top.q3_raw.iloc[5]:.2f} points. SWPW ({top.q3_raw.iloc[2]:.1f}%) and SRP ({top.q3_raw.iloc[3]:.1f}%) complete the list.")
para("Compilation", bold=True, size=11.5, space=2)
for t in [
    "Population: the 25 Western scorecard BAs reporting in Q3 2026 (the 27 on the published scorecard less WACM and WAUW, retired 1 April 2026).",
    "Periods: months and quarters by the BA's reporting day (EIA-930 'Data Date'), not the UTC timestamp.",
    "Scored hours: reported demand present and a day-ahead forecast above zero, less four exclusions: (1) hours in which EIA "
    "imputed demand ('Demand (MW) (Imputed)' populated); (2) forecast-flatline days, where every forecast the BA reported for the "
    "day carries one value (two or more values; NEVP 692 MW on 2 July 2026 in Q3); (3) bedding-in: the effective reporting day of "
    "each active forecast-change notice and the 13 days after it (withdrawn notices do not count; LDWP 12 to 25 September 2026 in "
    "Q3); (4) daylight-saving changeover days, the reporting days with 23 or 25 hours.",
    "Measure: sum of |reported demand - day-ahead forecast| over sum of reported demand, times 100; quarters computed directly over their hours; ranked before rounding.",
    f"Reproduction: this compilation returns all {len(hist)} published figures ({(hist.period_type == 'month').sum()} monthly, "
    f"{(hist.period_type == 'quarter').sum()} quarterly, Q3 2024 to Q2 2026) exactly to the published decimal, as FCS-STD-2 section 5 requires (Annex A). "
    "The bedding-in length is pinned by the archive: 13 days returns 789 figures, 14 returns all 799, 15 returns 783.",
]:
    bullet(t)
para("The bedding-in reading in the migration note", bold=True, size=11.5, space=2)
eom = json.load(open("eom_reading.json"))
para(f"Running bedding-in to the end of the effective month, as d.whitcombe remembers, returns {eom['match']} of the 799 published figures. "
     "It misses 12 monthly and 8 quarterly figures, all in periods that contain a forecast-change notice (for example PSCO March 2025, "
     "5.4 published against 4.5 compiled, and CISO May 2026, 11.8 against 12.2), so FCS-STD-2 section 5 bars it. On that reading LDWP's "
     "bedding-in runs from 12 to 30 September, LDWP falls to 7.67% and eighth place, and CISO at 8.40% takes fifth: CISO would go on Watch "
     "instead of LDWP, with PACE sixth at 8.22%. The formal letter would still go to PSCO.")
doc.add_picture("q3_chart.png", width=Inches(6.6))
para("Q3 2026 scorecard", bold=True, size=11.5, space=2)
t = doc.add_table(rows=1, cols=7); t.style = "Light Grid Accent 1"
for i, h in enumerate(["Rank", "BA", "Q3 %", "Jul %", "Aug %", "Sep %", "Watch"]):
    t.rows[0].cells[i].text = h
for r in out.itertuples():
    c = t.add_row().cells
    for i, v in enumerate([r.q3_rank, r.ba, f"{r.q3_forecast_error_pct:.1f}", f"{r.jul_2026_pct:.1f}", f"{r.aug_2026_pct:.1f}", f"{r.sep_2026_pct:.1f}",
                           "Letter + Watch" if r.formal_letter == "Y" else ("Watch" if r.forecast_watch == "Y" else "")]):
        c[i].text = str(v)
for row in t.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs: run.font.size = Pt(8)
doc.add_page_break()
para(f"Annex A. Published figures against this compilation ({len(hist)} figures, all differences 0.0)", bold=True, size=11, space=4)
rows = hist[["period", "ba", "forecast_error_pct", "compiled", "difference"]].values.tolist()
n = (len(rows) + 2) // 3
t = doc.add_table(rows=1, cols=15); t.style = "Table Grid"
for blk in range(3):
    for i, h in enumerate(["Period", "BA", "Pub.", "Comp.", "Diff."]):
        t.rows[0].cells[blk * 5 + i].text = h
for i in range(n):
    c = t.add_row().cells
    for blk in range(3):
        j = blk * n + i
        if j < len(rows):
            p_, ba, pub, comp, d = rows[j]
            for k2, v in enumerate([p_, ba, f"{pub:.1f}", f"{comp:.1f}", f"{d:.1f}"]):
                c[blk * 5 + k2].text = v
for row in t.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs: run.font.size = Pt(6)
doc.save("../golden/scorecard_q3_2026_note.docx")
print(out.head(7).to_string(index=False))
