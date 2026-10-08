"""Builds the analysts' workbook template (scenario artifact). It forms domestically sourced growth as total population
change less net international migration, which keeps the Census residual in, and books the remainder after natural change
as 'domestic migration'. The agreement counts only the components of change, so the residual should be left out."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "inputs", "rrg_workbook_template.xlsx")
wb = Workbook(); inp = wb.active; inp.title = "Inputs"
blue = PatternFill("solid", fgColor="DDEBF7"); b = Font(bold=True)
inp["A1"] = "RRG evaluation workbook - template (board analysts, March 2026). Scenario artifact authored for this exercise."
inp["A2"] = "Paste Vintage 2025 values from co-est2025-alldata.csv into the blue cells. Domestically sourced growth is the year's population change with net international migration taken out."
cols = ["Area", "POPESTIMATE2022", "POPESTIMATE2023", "POPESTIMATE2024", "POPESTIMATE2025",
        "NPOPCHG2023", "NPOPCHG2024", "NPOPCHG2025", "INTERNATIONALMIG2023", "INTERNATIONALMIG2024", "INTERNATIONALMIG2025",
        "NATURALCHG2023", "NATURALCHG2024", "NATURALCHG2025"]
for j, c in enumerate(cols, 1):
    inp.cell(4, j, c).font = b
inp["A5"] = "Dallas County (48113)"; inp["A6"] = "Texas (state row, SUMLEV 040)"
for r in (5, 6):
    for j in range(2, len(cols) + 1):
        inp.cell(r, j).fill = blue
inp.column_dimensions["A"].width = 30
# B..E pop22..pop25 ; F..H NPOPCHG23..25 ; I..K INTL23..25 ; L..N NAT23..25
ev = wb.create_sheet("Evaluation")
for j, c in enumerate(["Area", "Year", "Start-of-year residents", "Domestically sourced growth per 1,000",
                       "Natural change per 1,000", "Domestic migration per 1,000"], 1):
    ev.cell(1, j, c).font = b
years = [("Check year 2022-23", 0), ("Base year 2023-24", 1), ("Program year 2024-25", 2)]
pop, npc, intl, nat = "BCDE", "FGH", "IJK", "LMN"
r = 2
for area, row, rest in [("Dallas County", 5, False), ("Rest of Texas", 6, True)]:
    for lab, k in years:
        if rest:
            base = f"(Inputs!{pop[k]}6-Inputs!{pop[k]}5)"; tot = f"((Inputs!{npc[k]}6-Inputs!{npc[k]}5)-(Inputs!{intl[k]}6-Inputs!{intl[k]}5))"
            nc = f"(Inputs!{nat[k]}6-Inputs!{nat[k]}5)"
        else:
            base = f"Inputs!{pop[k]}5"; tot = f"(Inputs!{npc[k]}5-Inputs!{intl[k]}5)"; nc = f"Inputs!{nat[k]}5"
        ev.cell(r, 1, area); ev.cell(r, 2, lab); ev.cell(r, 3, f"={base}")
        ev.cell(r, 4, f"={tot}/C{r}*1000"); ev.cell(r, 5, f"={nc}/C{r}*1000"); ev.cell(r, 6, f"=D{r}-E{r}")
        r += 1
ev["A10"] = "Effect (program year vs base year)"; ev["B10"] = "=(D4-D3)-(D7-D6)"
ev["A11"] = "Parallel-movement check (base year vs check year)"; ev["B11"] = "=(D3-D2)-(D6-D5)"
ev["A12"] = "Bar"; ev["B12"] = 1.8; ev["A13"] = "Limit"; ev["B13"] = 1.5
ev["A14"] = "Decision"; ev["B14"] = '=IF(AND(ABS(B11)<=B13,B10>=B12),"Extend","Lapse")'
for c in ("A10", "A11", "A14"): ev[c].font = b
ev.column_dimensions["A"].width = 46; ev.column_dimensions["B"].width = 22
for c in "CDEF": ev.column_dimensions[c].width = 26
wb.save(OUT); print("saved")
