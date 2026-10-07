import json
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm

ev = pd.read_csv("evidence.csv"); f = json.load(open("facts.json"))
cat = {c["category"]: c for c in f["by_category"]}
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.6)


def para(text="", bold=False, size=None, space=4):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space); return p


def mixed(parts, space=4):
    p = doc.add_paragraph()
    for t, b in parts:
        r = p.add_run(t); r.bold = b
    p.paragraph_format.space_after = Pt(space); return p


def head(text):
    p = para(text, bold=True, size=11.5, space=2); p.paragraph_format.space_before = Pt(8)


def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet"); p.paragraph_format.space_after = Pt(2)


gwh = lambda m: f"{m / 1000:+,.1f} GWh"
para("BF-0712 backfill decision (DQ-3318)", bold=True, size=15, space=2)
para("To: data governance board    From: data owner, published demand series    Date: 6 October 2026", size=9.5, space=8)

head("Decision: run BF-0712 with 21 incidents removed")
mixed([(f"Run BF-0712 with {f['removed_n']} of its {f['n']} incidents removed. ", True),
       (f"The amended job backfills {f['approved_n']} incidents covering {f['approved_hours']} BA-hours and changes published "
        f"demand by {gwh(f['applied_mwh'])} ({f['applied_mwh']:,} MWh). As submitted, the job would have backfilled all "
        f"{f['n']} incidents ({f['hours']:,} BA-hours) and changed published demand by {gwh(f['submitted_mwh'])} "
        f"({f['submitted_mwh']:,} MWh). ", False),
       ("Running it as submitted would have deleted more than a terawatt-hour of load that was really served. Rejecting it outright "
        "would leave 17 incidents of demand values that are plainly wrong in the published series.", False)])
mixed([("Why the job cannot run as submitted. ", True),
       ("The job treats every large gap between reported demand and the day-ahead forecast as bad demand. In 21 incidents the "
        f"reported demand is right. {cat['forecast error']['n']} are forecast errors ({cat['forecast error']['hours']} BA-hours; the "
        f"job would have changed demand by {gwh(cat['forecast error']['mwh'])}), and {cat['real load']['n']} are real changes in load "
        f"({cat['real load']['hours']} BA-hours; {gwh(cat['real load']['mwh'])}). DQ-STD-007 section 2 rules out backfilling either.", False)])

head("How each incident was decided")
for b in [
    "A BA's reported demand is derived from its metered net generation minus its total interchange (EIA-930 documentation). "
    "So a demand value is wrong when one of its inputs is wrong, or when the demand field itself is broken.",
    "For every incident I compared reported demand, the day-ahead forecast, net generation and total interchange with the same "
    "BA's previous seven days. I also rebuilt the BA's total interchange from what each directly interconnected neighbour "
    "reported for its side of the tie (EIA-930 INTERCHANGE files; the BA's own figure is used only where a neighbour is silent).",
    "Backfill when demand is reported as zero or negative for a retail BA, or when one input to demand (generation or interchange) "
    "departs from the rest of the evidence: interchange contradicted by the neighbours, or generation collapsing while "
    "neighbour-confirmed interchange stays put. A real loss of local generation would be met by more imports, not by a matching fall in load.",
    "Leave as reported when demand reconciles with generation and neighbour-confirmed interchange and moves the way real load "
    "moves: storm outages with gradual restoration, cold or hot weather. Also leave it when the forecast is the outlier against "
    "its own history (stuck, collapsed, or running persistently low) while demand matches its prior week.",
    "EIA's imputation flags do not settle these incidents. EIA imputes demand only when it is blank, zero, negative or at least "
    f"1.5 times the BA's past maximum, and it imputed just {f['eia_imputed_in_approved']} of the {f['approved_hours']} hours backfilled here.",
]:
    bullet(b)

head("Incident by incident")
t = doc.add_table(rows=1, cols=6); t.style = "Light Grid Accent 1"
for i, h in enumerate(["Incident", "BA, first hour (UTC), hours", "Verdict", "Submitted / applied (MWh)", "Demand vs prior week; forecast; interchange own vs neighbours (MW, medians)", "What settles it"]):
    t.rows[0].cells[i].text = h
for _, r in ev.iterrows():
    row = t.add_row().cells
    ev_txt = (f"D {r.demand_med:,.0f} vs {r.demand_prior_week_med:,.0f}; DF {r.forecast_med:,.0f}"
              + (" (one value all incident)" if r.forecast_distinct == 1 else "")
              + f"; NG {r.ng_med:,.0f} vs {r.ng_prior_week_med:,.0f}; TI {r.ti_own_med:,.0f} / {r.ti_neighbours_med:,.0f}")
    vals = [r.incident_id[-2:], f"{r.ba}, {r.first_hour_end_utc}, {r.hours}h",
            "Backfill" if r.verdict == "backfill" else f"Leave ({r.category})",
            f"{r.submitted_change_mwh:,} / {r.applied_change_mwh:,}", ev_txt, r.reason]
    for i, v in enumerate(vals):
        row[i].text = v
for row in t.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs: run.font.size = Pt(7.5)

head("Why no other decision survives")
para("Running the job as submitted fails DQ-STD-007 section 2 on 21 incidents. Rejecting it fails the purpose of the standard on 17 "
     "incidents where the published demand is plainly wrong (whole days of zero load at IID and TIDC; LDWP interchange reported as "
     "zero while its neighbours report about 1,100 MW of imports into it). Removing incidents is the only option that leaves every "
     "reported value that matches the evidence and replaces every one that does not. The largest single item, SWPP on 16 April 2026 "
     "(-636,189 MWh submitted), is a broken forecast, not broken demand. The three closest calls are LDWP from 17 November 2024, "
     "WAUW from 17 December 2025 and SRP from 4 May 2026. All three are backfilled: in each one input to demand departs from the "
     "evidence while the others hold (two generation readings that collapse with interchange unchanged, and one interchange "
     "reading that its neighbours contradict).", space=4)
doc.save("../golden/BF-0712_decision_memo.docx")
print("saved")
